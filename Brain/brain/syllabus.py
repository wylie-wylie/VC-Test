from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import List, Optional

from brain.models import ScheduleRow


TABLE_HEADER = "| Date (YYYY-MM-DD) | Module | Topic | Deliverable |"
TABLE_DIVIDER = "| --- | --- | --- | --- |"

_DOW = (
    "Monday",
    "Tuesday",
    "Wednesday",
    "Thursday",
    "Friday",
    "Saturday",
    "Sunday",
    "Mon",
    "Tue",
    "Wed",
    "Thu",
    "Fri",
    "Sat",
    "Sun",
    "Th",
    "Sa",
    "Su",
    "M",
    "T",
    "W",
    "F",
)
_DOW_RE = "|".join(sorted(set(_DOW), key=len, reverse=True))
_GRID_LINE = re.compile(
    rf"^(?P<md>\d{{1,2}}/\d{{1,2}})\s+(?P<dow>{_DOW_RE})\b\s*(?P<rest>.*)$"
)
_MONTHS = {
    "january": 1,
    "february": 2,
    "march": 3,
    "april": 4,
    "may": 5,
    "june": 6,
    "july": 7,
    "august": 8,
    "september": 9,
    "october": 10,
    "november": 11,
    "december": 12,
}
_DOW_INDEX = {
    "monday": 0,
    "mon": 0,
    "m": 0,
    "tuesday": 1,
    "tue": 1,
    "t": 1,
    "wednesday": 2,
    "wed": 2,
    "w": 2,
    "thursday": 3,
    "thu": 3,
    "th": 3,
    "friday": 4,
    "fri": 4,
    "f": 4,
    "saturday": 5,
    "sat": 5,
    "sa": 5,
    "sunday": 6,
    "sun": 6,
    "su": 6,
}
_UNIT_WORD = {
    "one": "1",
    "two": "2",
    "three": "3",
    "four": "4",
    "five": "5",
    "six": "6",
}


def cell(value: str) -> str:
    return (value or "").replace("|", "/").replace("\n", " ").strip()


def render_table(rows: List[ScheduleRow]) -> str:
    lines = [
        TABLE_HEADER,
        TABLE_DIVIDER,
    ]
    for row in rows:
        when = row.date.isoformat() if row.date else ""
        lines.append(
            f"| {when} | {cell(row.module)} | {cell(row.topic)} | {cell(row.deliverable)} |"
        )
    return "\n".join(lines) + "\n"


def parse_table(text: str) -> List[ScheduleRow]:
    rows = []
    for line in text.splitlines():
        if not line.startswith("|"):
            continue
        parts = [part.strip() for part in line.strip().strip("|").split("|")]
        if len(parts) < 4 or parts[0].lower().startswith("date") or set(parts[0]) <= set("-: "):
            continue
        when = None
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", parts[0]):
            when = date.fromisoformat(parts[0])
        rows.append(ScheduleRow(when, parts[1], parts[2], parts[3]))
    return rows


def load_schedule(path: Path) -> List[ScheduleRow]:
    if not path.is_file():
        return []
    return parse_table(path.read_text())


def _year_for_month(month: int, term_year: int, term_start: Optional[date]) -> int:
    if term_start and term_start.month >= 7 and month <= 6:
        return term_year + 1
    return term_year


def _align_weekday(candidate: date, dow_token: str, previous: Optional[date]) -> date:
    target = _DOW_INDEX.get(dow_token.lower())
    if target is None or candidate.weekday() == target:
        if previous and candidate < previous:
            shifted = _shift_until_weekday(candidate, target, previous)
            return shifted
        return candidate
    return _shift_until_weekday(candidate, target, previous)


def _shift_until_weekday(candidate: date, target: Optional[int], previous: Optional[date]) -> date:
    if target is None:
        return candidate
    best = candidate
    best_distance = 999
    for delta_month in (-1, 0, 1, 2):
        month = candidate.month + delta_month
        year = candidate.year
        while month < 1:
            month += 12
            year -= 1
        while month > 12:
            month -= 12
            year += 1
        try:
            trial = date(year, month, candidate.day)
        except ValueError:
            continue
        if trial.weekday() != target:
            continue
        if previous and trial < previous:
            continue
        distance = abs((trial - candidate).days)
        if distance < best_distance:
            best = trial
            best_distance = distance
    return best


def _clean_topic(value: str) -> str:
    value = re.sub(r"\s+", " ", value).strip(" ,;|-")
    value = re.sub(r"(?:\s+CL|\s+Q|\s+HW)+$", "", value, flags=re.I).strip(" ,")
    trimmed = re.sub(r"\s+\d+(?:&\d+)?$", "", value).strip(" ,")
    if re.search(r"(?i)(exam|midterm|module|chapter)\s+\d+$", value):
        return value
    return trimmed


def _parse_grid_rest(rest: str, carried: str):
    rest = re.sub(r"\s+", " ", rest).strip()
    if re.search(r"\bno class\b", rest, re.I):
        return carried, rest, "No class", carried
    hard_exam = re.search(r"\b(Midterm\s*\d+|Final Exam)\b", rest, re.I)
    if hard_exam:
        phrase = re.sub(r"\s+", " ", hard_exam.group(1)).strip()
        return phrase, phrase, phrase, carried
    if re.search(r"\bExam Review\b", rest, re.I):
        topic = f"{carried} exam review".strip() if carried else "Exam Review"
        return carried or "Exam Review", topic, "Exam Review", carried
    lecture = ""
    lecture_match = re.search(r"\b(M\d+L\d+)\b", rest)
    if lecture_match:
        lecture = lecture_match.group(1)
    module = carried
    module_match = re.search(
        r"((?:Module\s+\d+|M\d+)\s*:\s*.+)",
        rest,
        re.I,
    )
    if module_match:
        module = _clean_topic(module_match.group(1))
        carried = module
    exam_match = re.search(
        r"((?:Midterm\s*\d+|Final Exam|Exam Review)(?:,\s*[^|]{0,40})?)",
        rest,
        re.I,
    )
    deliverable_parts = []
    topic = module
    if lecture:
        topic = f"{lecture} — {module}" if module else lecture
        deliverable_parts.append(lecture)
    if exam_match:
        phrase = _clean_topic(exam_match.group(1))
        phrase = re.sub(r",?\s+\d{1,2}:\d{2}.*$", "", phrase).strip(" ,")
        deliverable_parts.append(phrase)
        if re.search(r"midterm|final", phrase, re.I):
            topic = phrase
            module = phrase
    tail_source = rest
    if module_match:
        tail_source = rest[module_match.end() :]
    tails = re.findall(r"\b(\d+(?:&\d+)?|CL)\b", tail_source)
    if tails:
        deliverable_parts.append(" ".join(tails))
    deliverable = ", ".join(part for part in deliverable_parts if part)
    if not topic:
        topic = _clean_topic(rest)
    return module, topic, deliverable, carried


def parse_grid_text(text: str, term_year: int, term_start: Optional[date] = None) -> List[ScheduleRow]:
    rows: List[ScheduleRow] = []
    carried = ""
    previous: Optional[date] = None
    for raw_line in text.splitlines():
        line = raw_line.strip()
        match = _GRID_LINE.match(line)
        if not match:
            final = re.search(
                r"Final Exam:\s*([A-Za-z]+),?\s*([A-Za-z]+)\s+(\d{1,2})",
                line,
                re.I,
            )
            if final and final.group(2).lower() in _MONTHS:
                month = _MONTHS[final.group(2).lower()]
                year = _year_for_month(month, term_year, term_start)
                when = _align_weekday(
                    date(year, month, int(final.group(3))),
                    final.group(1),
                    previous,
                )
                rows.append(ScheduleRow(when, "Final Exam", "Final Exam", "Final Exam"))
                previous = when
            continue
        month_s, day_s = match.group("md").split("/")
        month, day = int(month_s), int(day_s)
        year = _year_for_month(month, term_year, term_start)
        try:
            candidate = date(year, month, day)
        except ValueError:
            continue
        when = _align_weekday(candidate, match.group("dow"), previous)
        module, topic, deliverable, carried = _parse_grid_rest(match.group("rest"), carried)
        rows.append(ScheduleRow(when, module, topic, deliverable))
        previous = when
    return rows


def _parse_named_date(month_name: str, day: str, dow: str, term_year: int, term_start: Optional[date]) -> date:
    month = _MONTHS[month_name.lower()]
    year = _year_for_month(month, term_year, term_start)
    return _align_weekday(date(year, month, int(day)), dow, None)


def parse_narrative_text(text: str, term_year: int, term_start: Optional[date] = None) -> List[ScheduleRow]:
    units = []
    current = None
    for line in text.splitlines():
        unit_match = re.search(r"Unit\s+([A-Za-z0-9]+)\s+Content", line, re.I)
        if unit_match:
            label = unit_match.group(1)
            label = _UNIT_WORD.get(label.lower(), label)
            current = {"module": f"Unit {label}", "chapters": []}
            units.append(current)
            continue
        chapter = re.search(r"Chapter\s+(\d+)\s*:\s*(.+)", line)
        if chapter and current is not None:
            title = _clean_topic(chapter.group(2))
            current["chapters"].append(f"Ch {chapter.group(1)}: {title}")
    exams = []
    for match in re.finditer(
        r"(Exam\s*\d+|Final Exam)\s*:\s*(?:\([^)]*\)\s*)?([A-Za-z]+),?\s*([A-Za-z]+)\s+(\d{1,2})",
        text,
        re.I,
    ):
        if match.group(3).lower() not in _MONTHS:
            continue
        when = _parse_named_date(match.group(3), match.group(4), match.group(2), term_year, term_start)
        exams.append((when, _clean_topic(match.group(1))))
    finals = []
    for match in re.finditer(
        r"Sections?\s+[0-9][0-9a-z,\s]*:\s*([A-Za-z]+)\s+([A-Za-z]+)\s+(\d{1,2})",
        text,
        re.I,
    ):
        if match.group(2).lower() not in _MONTHS:
            continue
        finals.append(
            _parse_named_date(match.group(2), match.group(3), match.group(1), term_year, term_start)
        )
    if finals and not any(when.month == 12 for when, _label in exams):
        exams.append((min(finals), "Final Exam"))
    exams.sort()
    rows: List[ScheduleRow] = []
    cursor = term_start
    for index, (when, label) in enumerate(exams):
        unit = units[index] if index < len(units) else None
        module = unit["module"] if unit else label
        topic = "; ".join(unit["chapters"]) if unit and unit["chapters"] else label
        if cursor and (not rows or rows[-1].date != cursor):
            rows.append(ScheduleRow(cursor, module, topic, "Workshop sequence"))
        elif cursor is None and unit:
            rows.append(ScheduleRow(when, module, topic, label))
        rows.append(ScheduleRow(when, module, topic, label))
        cursor = when.fromordinal(when.toordinal() + 1)
    if units and len(units) > len(exams) and cursor:
        unit = units[len(exams)]
        rows.append(
            ScheduleRow(
                cursor,
                unit["module"],
                "; ".join(unit["chapters"]),
                "Workshop sequence",
            )
        )
    deduped = []
    seen = set()
    for row in rows:
        key = (row.date, row.module, row.deliverable)
        if key in seen:
            continue
        seen.add(key)
        deduped.append(row)
    return deduped


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        chunks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(chunks)
    return path.read_text(errors="replace")


def convert_schedule(path: Path, term_year: int, term_start: Optional[date] = None) -> str:
    if path.suffix.lower() in {".md", ".markdown", ".txt"} and TABLE_HEADER in path.read_text():
        rows = parse_table(path.read_text())
        return render_table(rows)
    text = extract_text(path)
    grid_hits = sum(1 for line in text.splitlines() if _GRID_LINE.match(line.strip()))
    if grid_hits >= 5:
        rows = parse_grid_text(text, term_year, term_start)
    else:
        rows = parse_narrative_text(text, term_year, term_start)
    if not rows:
        rows = []
    return render_table(rows)


def lookup_row(rows: List[ScheduleRow], day: date) -> Optional[ScheduleRow]:
    exact = [row for row in rows if row.date == day]
    if exact:
        return exact[-1]
    prior = [row for row in rows if row.date and row.date < day and not row.is_exam and not row.is_no_class]
    if not prior:
        return None
    return max(prior, key=lambda row: row.date)


def topic_label(row: Optional[ScheduleRow], fallback: str) -> str:
    if row is None:
        return fallback
    if row.is_no_class:
        return row.topic or "No class"
    topic = row.topic or row.module or fallback
    if row.deliverable and row.deliverable not in topic and not row.deliverable.startswith("M"):
        if row.is_exam:
            return row.deliverable
    return topic


def newest_slide(slides_dir: Path) -> Optional[Path]:
    if not slides_dir.is_dir():
        return None
    files = [
        path
        for path in slides_dir.iterdir()
        if path.is_file() and path.suffix.lower() in {".pdf", ".md", ".markdown"} and not path.name.startswith(".")
    ]
    if not files:
        return None
    return max(files, key=lambda path: path.stat().st_mtime)


def slide_blurb(path: Optional[Path]) -> str:
    if path is None:
        return ""
    if path.suffix.lower() in {".md", ".markdown"}:
        for line in path.read_text(errors="replace").splitlines():
            stripped = line.strip().lstrip("#").strip()
            if stripped:
                return stripped[:180]
        return path.stem
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        text = (reader.pages[0].extract_text() or "") if reader.pages else ""
    except Exception:
        return path.stem
    lines = [re.sub(r"\s+", " ", line).strip(" ●•-\t") for line in text.splitlines()]
    lines = [line for line in lines if line and not re.fullmatch(r"lecture\s*\d+", line, re.I)]
    if not lines:
        return path.stem
    return "; ".join(lines[:6])[:180]


def week_rows(rows: List[ScheduleRow], start: date, end: date) -> List[ScheduleRow]:
    return [row for row in rows if row.date and start <= row.date <= end and not row.is_no_class]
